from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# Palette
_DARK = RGBColor(0x1A, 0x1A, 0x2E)
_MID = RGBColor(0x47, 0x55, 0x69)
_LIGHT = RGBColor(0x64, 0x74, 0x8B)
_RULE = "94A3B8"


def _section_header(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = _DARK

    # Draw a thin rule under the header using OOXML directly — python-docx
    # doesn't expose paragraph borders through its high-level API.
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), _RULE)
    pBdr.append(bottom)
    pPr.append(pBdr)


def write_resume(data: dict, path: str) -> None:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    doc.styles["Normal"].paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(data.get("name", ""))
    r.bold = True
    r.font.size = Pt(20)

    if data.get("contact"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(data["contact"])
        r.font.size = Pt(10)
        r.font.color.rgb = _MID

    if data.get("summary"):
        _section_header(doc, "Summary")
        p = doc.add_paragraph(data["summary"])
        p.paragraph_format.space_after = Pt(2)
        p.runs[0].font.size = Pt(10)

    if data.get("experience"):
        _section_header(doc, "Experience")
        for exp in data["experience"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(exp.get("company", ""))
            r.bold = True
            r.font.size = Pt(10)
            if exp.get("dates"):
                r = p.add_run(f"   {exp['dates']}")
                r.font.size = Pt(10)
                r.font.color.rgb = _LIGHT

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(exp.get("title", ""))
            r.italic = True
            r.font.size = Pt(10)
            r.font.color.rgb = _MID

            for bullet in exp.get("bullets", []):
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_after = Pt(1)
                p.add_run(bullet).font.size = Pt(10)

    if data.get("skills"):
        _section_header(doc, "Skills")
        p = doc.add_paragraph(", ".join(data["skills"]))
        p.paragraph_format.space_after = Pt(2)
        p.runs[0].font.size = Pt(10)

    if data.get("education"):
        _section_header(doc, "Education")
        for edu in data["education"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(edu.get("school", ""))
            r.bold = True
            r.font.size = Pt(10)
            parts = [x for x in [edu.get("degree"), edu.get("year")] if x]
            if parts:
                r = p.add_run(f"   {' · '.join(parts)}")
                r.font.size = Pt(10)
                r.font.color.rgb = _LIGHT

    doc.save(path)


def write_cover_letter(text: str, path: str) -> None:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

    doc.styles["Normal"].paragraph_format.space_after = Pt(0)

    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        # Collapse single newlines (e.g. inside the sign-off block) into spaces
        # so they render as a single paragraph rather than line breaks.
        p.add_run(block.replace("\n", " ")).font.size = Pt(11)

    doc.save(path)
